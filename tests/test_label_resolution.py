"""
test_label_resolution.py

Covers models.Selector.from_input (label vs. cell auto-detection) and
spreadsheet_processor.resolve_selector (finding a selector's value inside
a worksheet), including the "Handle Inconsistent Spreadsheet Layouts"
design requirement: a label must still be found after it moves to a
different row or a different starting cell.
"""

import tempfile
import unittest
from pathlib import Path

import openpyxl

from models import Selector, SelectorType
from spreadsheet_processor import load_workbook_safe, resolve_selector, select_worksheet

TEST_DATA_DIR = Path(__file__).resolve().parent.parent / "test_data"


class TestSelectorFromInput(unittest.TestCase):
    def test_plain_word_is_treated_as_a_label(self):
        selector = Selector.from_input("TEMP")
        self.assertEqual(selector.type, SelectorType.LABEL)

    def test_dollar_ticker_is_treated_as_a_label(self):
        selector = Selector.from_input("$AAPL")
        self.assertEqual(selector.type, SelectorType.LABEL)

    def test_standard_cell_reference_is_detected(self):
        for raw in ("A1", "b12", "Z9", "AA123"):
            with self.subTest(raw=raw):
                self.assertEqual(Selector.from_input(raw).type, SelectorType.CELL)

    def test_series_name_defaults_to_raw_input(self):
        selector = Selector.from_input("  HUMIDITY  ")
        self.assertEqual(selector.raw, "HUMIDITY")
        self.assertEqual(selector.series_name, "HUMIDITY")

    def test_empty_selector_raises(self):
        with self.assertRaises(ValueError):
            Selector(raw="   ", type=SelectorType.LABEL, series_name="")


class TestResolveSelector(unittest.TestCase):
    def _worksheet(self, relative_path: str):
        workbook, error = load_workbook_safe(TEST_DATA_DIR / relative_path)
        self.assertIsNone(error, msg=f"failed to open {relative_path}: {error}")
        self.addCleanup(workbook.close)
        worksheet, sheet_error = select_worksheet(workbook, sheet_name=None)
        self.assertIsNone(sheet_error)
        return worksheet

    def test_label_selector_finds_value_in_standard_layout(self):
        worksheet = self._worksheet("valid/weather_01.xlsx")
        value, cell, note = resolve_selector(worksheet, Selector.from_input("TEMP"))
        self.assertEqual(value, 72)
        self.assertEqual(cell, "B2")
        self.assertIsNone(note)

    def test_cell_selector_reads_the_exact_address(self):
        worksheet = self._worksheet("valid/weather_01.xlsx")
        value, cell, note = resolve_selector(worksheet, Selector.from_input("B2"))
        self.assertEqual(value, 72)
        self.assertEqual(cell, "B2")
        self.assertIsNone(note)

    def test_label_survives_row_reordering(self):
        # shifted_01.xlsx puts HUMIDITY before TEMP, unlike the standard
        # fixtures - resolution must search for the label, not assume a
        # fixed row.
        worksheet = self._worksheet("shifted_labels/shifted_01.xlsx")
        value, _cell, _note = resolve_selector(worksheet, Selector.from_input("TEMP"))
        self.assertEqual(value, 72)

    def test_label_survives_table_moving_to_a_new_starting_cell(self):
        # shifted_02.xlsx moves the whole Property/Value table to C3:D5.
        worksheet = self._worksheet("shifted_labels/shifted_02.xlsx")
        value, cell, _note = resolve_selector(worksheet, Selector.from_input("TEMP"))
        self.assertEqual(value, 74)
        self.assertEqual(cell, "D4")

    def test_missing_label_reports_error_instead_of_raising(self):
        worksheet = self._worksheet("missing_values/weather_missing_02.xlsx")
        value, cell, note = resolve_selector(worksheet, Selector.from_input("TEMP"))
        self.assertIsNone(value)
        self.assertIsNone(cell)
        self.assertIn("not found", note)

    def test_extra_columns_do_not_confuse_resolution(self):
        worksheet = self._worksheet("mixed_structure/extra_columns.xlsx")
        value, _cell, _note = resolve_selector(worksheet, Selector.from_input("HUMIDITY"))
        self.assertEqual(value, 52)

    def test_duplicate_label_resolves_to_the_first_match_without_error(self):
        # Build a workbook with TEMP appearing twice, on purpose. The plan
        # only asks that a moved label still be found, not that duplicates
        # be specially detected - this just confirms the first match wins
        # cleanly rather than raising or behaving unpredictably.
        with tempfile.TemporaryDirectory() as tmp:
            path = Path(tmp) / "duplicate_label.xlsx"
            workbook = openpyxl.Workbook()
            sheet = workbook.active
            sheet.append(["Property", "Value"])
            sheet.append(["TEMP", 72])
            sheet.append(["TEMP", 99])
            workbook.save(path)

            loaded, error = load_workbook_safe(path)
            self.assertIsNone(error)
            self.addCleanup(loaded.close)
            worksheet, _ = select_worksheet(loaded, sheet_name=None)

            value, cell, error = resolve_selector(worksheet, Selector.from_input("TEMP"))

            self.assertEqual(value, 72)
            self.assertEqual(cell, "B2")
            self.assertIsNone(error)


if __name__ == "__main__":
    unittest.main()
